"""A month of sales for the accountant (Phase 15, D72/D73): an Excel file.

Owners pick a month in the Mini App (Orders → Export), Gregorian (GC) or
Ethiopian (ዓ.ም., Phase 15c, D79); the shop's bot sends the file to their
private chat (files opened inside Telegram's Mini App often don't download).
Every date is shown in both calendars. Five sheets:

    Summary     money received, per channel and payment method, discounts,
                unpaid, average sale, delivery / pickup / in the shop
    By day      each day: orders, paid, money received, discounts (the cash book)
    By product  what sold: quantity, money, discounts, best first
    Sales       one row per item of every paid order: what, how many, listed and
                final price, discount, payment method, who sold or confirmed it
    Orders      one row per order, paid or not: customer, delivery/pickup, stage

Months are Addis Ababa months, by the day the order was placed. An Ethiopian
Nehase includes Pagume (the 5–6 days before the new year), so the year has 12
exports. Cancelled orders are listed in Orders but count nowhere else.
"""
import io
import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from typing import Any, Literal
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.agents.ethiopian import MONTHS_AM, MONTHS_EN, ethiopian_text, month_range
from app.agents.shop_types import labels
from app.agents.tools import order_number
from app.models.schemas import Store

ADDIS = timezone(timedelta(hours=3), "EAT")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MONEY = "#,##0.00"
DATE = "yyyy-mm-dd hh:mm"
DAY = "yyyy-mm-dd"

Calendar = Literal["gregorian", "ethiopian"]
STAGES = {"pending": "Waiting for payment", "confirmed": "Paid", "out_for_delivery": "On the way",
          "delivered": "Delivered", "cancelled": "Cancelled"}
CHANNELS = {"telegram": "Telegram", "in_shop": "In the shop"}
_MONTH = re.compile(r"^(\d{4})-(\d{2})$")


class ExportError(Exception):
    def __init__(self, message: str):
        super().__init__(message)
        self.message = message


@dataclass
class Period:
    """The month asked for: [start, end) in Addis Ababa, and its names."""
    start: datetime
    end: datetime
    title: str  # "September 2026" / "መስከረም 2019 ዓ.ም."
    other: str  # the same days in the other calendar
    slug: str  # for the file name: "2026-09" / "Meskerem-2019-EC"


def _gc_day(value: date) -> str:
    return value.strftime("%b %d, %Y").replace(" 0", " ")  # "Sep 11, 2026"


def month_bounds(month: str, now: datetime | None = None, calendar: Calendar = "gregorian") -> Period:
    """"2026-09" (GC) or "2019-01" (ዓ.ም.: Meskerem 2019) → its days in Addis
    Ababa. Not a month that hasn't started yet."""
    match = _MONTH.match(month or "")
    year, number = (int(match[1]), int(match[2])) if match else (0, 0)
    if calendar == "ethiopian":
        if not 1 <= number <= 12 or year < 2012:
            raise ExportError("Choose an Ethiopian month like 2019-01 (Meskerem 2019).")
        first, after = month_range(year, number)
        start = datetime(first.year, first.month, first.day, tzinfo=ADDIS)
        end = datetime(after.year, after.month, after.day, tzinfo=ADDIS)
        name = MONTHS_AM[number - 1] + (f" + {MONTHS_AM[12]}" if number == 12 else "")
        period = Period(start, end, f"{name} {year} ዓ.ም. ({MONTHS_EN[number - 1]})",
                        f"{_gc_day(first)} – {_gc_day(after - timedelta(days=1))}",
                        f"{MONTHS_EN[number - 1]}-{year}-EC")
    else:
        if not 1 <= number <= 12 or year < 2020:
            raise ExportError("Choose a month like 2026-09.")
        start = datetime(year, number, 1, tzinfo=ADDIS)
        end = datetime(year + (number == 12), number % 12 + 1, 1, tzinfo=ADDIS)
        last = (end - timedelta(days=1)).date()
        period = Period(start, end, start.strftime("%B %Y"),
                        f"{ethiopian_text(start.date())} – {ethiopian_text(last)} ዓ.ም.", f"{month}")
    if period.start > (now or datetime.now(timezone.utc)):
        raise ExportError("That month hasn't started yet.")
    return period


@dataclass
class Export:
    filename: str
    data: bytes
    orders: int
    revenue: Decimal


def _money(value: Any) -> Decimal:
    return Decimal(str(value or 0))


def _local(value: str) -> datetime:
    """A database time as Addis local time, without a zone (Excel has none)."""
    return datetime.fromisoformat(value).astimezone(ADDIS).replace(tzinfo=None)


def _paid_amount(order: dict[str, Any]) -> Decimal:
    payments = order.get("payments") or []
    return sum((_money(p.get("amount")) for p in payments), Decimal(0)) if payments \
        else _money(order.get("total_price"))


def _method(order: dict[str, Any]) -> str:
    payments = order.get("payments") or []
    return order.get("payment_method") or next((p["method"] for p in payments if p.get("method")), "") or ""


def _who(order: dict[str, Any]) -> str:
    payments = order.get("payments") or []
    return order.get("sold_by_name") or next(
        (p["confirmed_by_name"] for p in payments if p.get("confirmed_by_name")), "") or ""


def _sheet(book: Workbook, title: str, header: list[str], rows: list[list[Any]],
           money: tuple[int, ...] = (), dates: tuple[int, ...] = (), days: tuple[int, ...] = ()) -> None:
    sheet = book.create_sheet(title)
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for column in range(1, len(header) + 1):
        letter = get_column_letter(column)
        for cell in sheet[letter][1:]:
            if column in money:
                cell.number_format = MONEY
            elif column in dates:
                cell.number_format = DATE
            elif column in days:
                cell.number_format = DAY
        widest = max([len(str(header[column - 1]))] + [len(str(r[column - 1] or "")) for r in rows])
        sheet.column_dimensions[letter].width = min(max(widest + 2, 10), 45)
    sheet.freeze_panes = "A2"


@dataclass
class _Day:
    placed: int = 0
    paid: int = 0
    money: Decimal = Decimal(0)
    discounts: Decimal = Decimal(0)


def build_export(store: Store, orders: list[dict[str, Any]], period: Period,
                 now: datetime | None = None) -> Export:
    """The workbook for `orders` (from orders_for_export), as bytes."""
    option1, option2 = labels(store)
    sales: list[list[Any]] = []
    order_rows: list[list[Any]] = []
    revenue = discounts = Decimal(0)
    by_channel: dict[str, Decimal] = {}
    by_method: dict[str, tuple[int, Decimal]] = {}
    by_product: dict[tuple[str, str], list[Any]] = {}  # (name, code) -> [quantity, money, discounts]
    kinds = {"delivery": 0, "pickup": 0, "in_shop": 0}
    today = (now or datetime.now(timezone.utc)).astimezone(ADDIS).date()
    last_day = min(period.end.date() - timedelta(days=1), today)
    days = {period.start.date() + timedelta(days=n): _Day()
            for n in range(max((last_day - period.start.date()).days + 1, 0))}
    paid_count = unpaid_count = cancelled_count = 0

    for order in orders:
        number = order_number(UUID(order["id"]))
        when = _local(order["created_at"])
        ethiopian = ethiopian_text(when.date())
        channel_key = order.get("channel") or "telegram"
        channel = CHANNELS.get(channel_key, "Telegram")
        status, paid = order.get("status") or "pending", order.get("payment_status") == "paid"
        items = order.get("order_items") or []
        order_rows.append([
            when, ethiopian, number, channel, order.get("contact_name") or "", order.get("contact_phone") or "",
            (order.get("fulfillment_method") or "").capitalize(), STAGES.get(status, status),
            "Paid" if paid else ("Refunded" if order.get("payment_status") == "refunded" else "Not paid"),
            _money(order.get("total_price")), _method(order) if paid else "", sum(i["quantity"] for i in items),
        ])
        if status == "cancelled":
            cancelled_count += 1
            continue
        day = days.setdefault(when.date(), _Day())
        day.placed += 1
        kinds["in_shop" if channel_key == "in_shop" else order.get("fulfillment_method") or "pickup"] += 1
        if not paid:
            unpaid_count += 1
            continue
        paid_count += 1
        amount = _paid_amount(order)
        revenue += amount
        day.paid += 1
        day.money += amount
        by_channel[channel] = by_channel.get(channel, Decimal(0)) + amount
        method = _method(order) or "Not recorded"
        count, total = by_method.get(method, (0, Decimal(0)))
        by_method[method] = (count + 1, total + amount)
        for item in items:
            variant = item.get("product_variants") or {}
            product = variant.get("products") or {}
            price, quantity = _money(item.get("price")), int(item["quantity"])
            listed = _money(item["list_price"]) if item.get("list_price") is not None else price
            off = max(listed - price, Decimal(0)) * quantity
            discounts += off
            day.discounts += off
            line = by_product.setdefault((product.get("name") or "", product.get("code") or ""),
                                         [0, Decimal(0), Decimal(0)])
            line[0] += quantity
            line[1] += price * quantity
            line[2] += off
            sales.append([when, ethiopian, number, channel, product.get("name") or "", product.get("code") or "",
                          variant.get("color") or "", variant.get("size") or "", quantity, listed, price,
                          off, price * quantity, _method(order), _who(order)])

    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    summary_rows: list[list[Any]] = [
        [store.name, period.title],
        ["", period.other],
        [],
        ["Money received (paid orders and sales)", revenue],
        ["Paid orders and sales", paid_count],
        *[[f"  {channel}", total] for channel, total in sorted(by_channel.items())],
        ["Average sale", (revenue / paid_count).quantize(Decimal("0.01")) if paid_count else Decimal(0)],
        ["Discounts given in the shop", discounts],
        ["Orders not paid", unpaid_count],
        ["Cancelled orders", cancelled_count],
        [],
        ["Orders by kind", "Orders"],
        ["  Delivery (Telegram)", kinds["delivery"]],
        ["  Pickup (Telegram)", kinds["pickup"]],
        ["  In the shop", kinds["in_shop"]],
        [],
        ["By payment method", "Payments", "Amount"],
        *[[method, count, total] for method, (count, total) in sorted(by_method.items())],
        [],
        ["Orders placed in this month, Addis Ababa time. Amounts in ETB."],
        [f"Made {today:%Y-%m-%d} ({ethiopian_text(today)} ዓ.ም.)"],
    ]
    for row in summary_rows:
        summary.append(row)
    summary["A1"].font = Font(bold=True, size=14)
    summary["B1"].font = Font(bold=True, size=14)
    for row in summary.iter_rows():
        for cell in row:
            if isinstance(cell.value, Decimal):
                cell.number_format = MONEY
            if cell.value in ("Orders by kind", "By payment method", "Orders", "Payments", "Amount"):
                cell.font = Font(bold=True)
    summary.column_dimensions["A"].width = 40
    summary.column_dimensions["B"].width = 36
    summary.column_dimensions["C"].width = 16

    _sheet(book, "By day",
           ["Date", "Ethiopian date", "Orders", "Paid", "Money received", "Discounts"],
           [[day, ethiopian_text(day), d.placed, d.paid, d.money, d.discounts] for day, d in sorted(days.items())],
           money=(5, 6), days=(1,))
    _sheet(book, "By product",
           ["Product", "Code", "Quantity sold", "Money", "Discounts"],
           [[name, code, q, m, off] for (name, code), (q, m, off)
            in sorted(by_product.items(), key=lambda item: (-item[1][1], item[0]))],
           money=(4, 5))
    _sheet(book, "Sales",
           ["Date", "Ethiopian date", "Order", "Where", "Product", "Code", option1.en, option2.en, "Quantity",
            "Listed price", "Price", "Discount", "Total", "Payment method", "Sold / confirmed by"],
           sales, money=(10, 11, 12, 13), dates=(1,))
    _sheet(book, "Orders",
           ["Date", "Ethiopian date", "Order", "Where", "Customer", "Phone", "Delivery / pickup", "Stage",
            "Payment", "Total", "Payment method", "Items"],
           order_rows, money=(10,), dates=(1,))

    buffer = io.BytesIO()
    book.save(buffer)
    shop = re.sub(r"[^A-Za-z0-9]+", "-", store.name).strip("-") or "shop"
    return Export(f"{shop}-{period.slug}.xlsx", buffer.getvalue(), len(orders), revenue)
