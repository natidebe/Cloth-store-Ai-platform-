"""Phase 15, quick wins (D70–D75) and 15b (D76–D78): the owner's morning
summary, the month's export for the accountant, and order updates to
customers from the staff group (delivery: paid on arrival). No network: Telegram and the database are faked."""
import io
import json
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from uuid import uuid4

import httpx
import pytest
from openpyxl import load_workbook

from app.agents.daily_summary import ADDIS, DailySummaries, summary_text, yesterday
from app.agents.ethiopian import ethiopian_text, from_ethiopian, to_ethiopian
from app.agents.export import ExportError, build_export, month_bounds
from app.agents.messages import t
from app.agents.staff import order_buttons
from app.agents.tools import order_number
from app.models.schemas import OrderWithItems, Store
from app.services.telegram_service import TELEGRAM_API, TelegramService
from tests import test_miniapp as mini
from tests.test_flow import (
    AF1_WHITE_42,
    CUSTOMER,
    PAYMENT_TEXT,
    STAFF_CHAT,
    STORE,
    World,
    placed_order,
    up_to_confirm,
)
from tests.test_miniapp import MEMBER, OWNER, STORE_A, world  # noqa: F401  (fixture)
from tests.test_staff import STAFF_MEMBER, _button_answers, _press, _staff

TOKEN = "111111111:AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA"
MORNING = datetime(2026, 10, 4, 9, 0, tzinfo=ADDIS)


@pytest.fixture
def anyio_backend():
    return "asyncio"


# --- 3. Order updates to customers (D74–D78) ------------------------------------------------

async def _paid_pickup(w: World):
    """A pickup order, paid with the staff group's button. Returns the order
    and the id of the staff note carrying its next steps."""
    order = await placed_order(w)
    await w.send_photo()
    await _staff(w, _press(f"pay:{order.id}", w.telegram.last_message_id(STAFF_CHAT)))
    return order, w.telegram.last_message_id(STAFF_CHAT)


async def _delivery_order(w: World):
    """A delivery order (paid on arrival, D29): the order and its staff alert."""
    await up_to_confirm(w, fulfillment="btn_delivery")
    await w.tap(t("btn_confirm", "en"))
    [order] = w.db.orders.values()
    return order, w.telegram.last_message_id(STAFF_CHAT)


def _order(w, order):
    return next(o for o in w.db.orders.values() if o.id == order.id)


def _buttons_now(w):
    """The buttons the last editMessageReplyMarkup left."""
    body = next(b for m, b, _ in reversed(w.telegram.calls) if m == "editMessageReplyMarkup")
    return [(b["text"], b["callback_data"]) for row in body["reply_markup"]["inline_keyboard"] for b in row]


@pytest.mark.anyio
async def test_a_delivery_order_goes_on_the_way_then_is_delivered_and_paid():
    w = World()
    order, alert = await _delivery_order(w)
    number = order_number(order.id)
    # D76: no "Confirm payment" first: the customer pays on arrival.
    assert w.telegram.buttons(alert) == [(f"🚚 On the way #{number}", f"ship:{order.id}"),
                                         ("▶️ Hand back to bot", f"resume:{CUSTOMER}")]
    assert w.db.variants[AF1_WHITE_42].stock_quantity == 2

    await _staff(w, _press(f"ship:{order.id}", alert))
    assert _order(w, order).status == "out_for_delivery"
    assert w.db.variants[AF1_WHITE_42].stock_quantity == 1  # D77: it left the shop
    told = w.telegram.to(CUSTOMER)[-1]
    assert told.startswith(t("order_on_the_way_pay", "en", number=number, total="5,000 ETB"))
    assert PAYMENT_TEXT in told and t("screenshot_on_delivery", "en") in told  # D78
    assert _buttons_now(w) == [(f"✅ Delivered & paid #{number}", f"done:{order.id}"),
                               (f"❌ Not delivered #{number}", f"back:{order.id}")]
    assert w.db.status_changes[-1] == (order.id, "out_for_delivery", STAFF_MEMBER, "Sara")  # who (D74)

    await _staff(w, _press(f"ship:{order.id}", alert))  # tapped twice
    assert _button_answers(w)[-1] == (f"Order #{number} is already on the way.", True)
    assert w.db.variants[AF1_WHITE_42].stock_quantity == 1  # not taken twice

    # The customer pays by transfer and sends the screenshot: staff see the order's next step.
    await w.send_photo(caption="paid")
    screenshot = w.telegram.last_message_id(STAFF_CHAT)
    assert (f"✅ Delivered & paid #{number}", f"done:{order.id}") in w.telegram.buttons(screenshot)

    await _staff(w, _press(f"done:{order.id}", screenshot))
    paid = _order(w, order)
    assert (paid.status, paid.payment_status) == ("delivered", "paid")
    assert w.db.payments[-1]["amount"] == order.total_price
    assert w.db.payments[-1]["confirmed_by"] == (STAFF_MEMBER, "Sara")
    assert w.db.variants[AF1_WHITE_42].stock_quantity == 1  # already taken at On the way
    assert w.telegram.to(CUSTOMER)[-1] == t("order_delivered_paid", "en", number=number, shop=STORE.name)
    assert _buttons_now(w) == []
    assert not w.conversation.bot_paused  # finished: the bot answers again

    await _staff(w, _press(f"done:{order.id}", alert))
    assert _button_answers(w)[-1] == (f"Order #{number} was already delivered.", True)
    assert len(w.db.payments) == 1


@pytest.mark.anyio
async def test_not_delivered_puts_the_stock_back_and_cancels():
    w = World()
    order, alert = await _delivery_order(w)
    number = order_number(order.id)
    await _staff(w, _press(f"ship:{order.id}", alert))
    assert w.db.variants[AF1_WHITE_42].stock_quantity == 1

    await _staff(w, _press(f"back:{order.id}", alert))
    assert _order(w, order).status == "cancelled"
    assert w.db.variants[AF1_WHITE_42].stock_quantity == 2
    assert w.telegram.to(CUSTOMER)[-1] == t("order_not_delivered", "en", number=number)
    assert _buttons_now(w) == [] and w.db.payments == []


@pytest.mark.anyio
async def test_delivered_needs_on_the_way_first_and_sold_out_is_refused():
    w = World()
    order, alert = await _delivery_order(w)
    number = order_number(order.id)
    await _staff(w, _press(f"done:{order.id}", alert))
    assert _button_answers(w)[-1] == (f"Tap 🚚 On the way for order #{number} first.", True)

    w.db.variants[AF1_WHITE_42] = w.db.variants[AF1_WHITE_42].model_copy(
        update={"stock_quantity": 0, "held": 0})  # sold in the shop meanwhile
    told = len(w.telegram.to(CUSTOMER))
    await _staff(w, _press(f"ship:{order.id}", alert))
    answer, popup = _button_answers(w)[-1]
    assert popup and "sold out" in answer and "Nothing was changed" in answer
    assert _order(w, order).status == "pending" and len(w.telegram.to(CUSTOMER)) == told


@pytest.mark.anyio
async def test_pickup_is_paid_first_then_ready_then_picked_up():
    w = World()
    order, note = await _paid_pickup(w)
    number = order_number(order.id)
    assert w.telegram.buttons(note) == [(f"📦 Ready for pickup #{number}", f"ready:{order.id}"),
                                        (f"✅ Picked up #{number}", f"done:{order.id}")]
    assert "Confirmed by Sara" in w.telegram.to(STAFF_CHAT)[-1]

    await _staff(w, _press(f"ready:{order.id}", note))
    assert _order(w, order).status == "confirmed"  # D75: a message only
    ready = w.telegram.to(CUSTOMER)[-1]
    assert ready.startswith(t("order_ready", "en", number=number))
    assert t("pickup_hours", "en", hours=STORE.opening_hours) in ready  # when to come, from the profile
    assert _buttons_now(w) == [(f"✅ Picked up #{number}", f"done:{order.id}")]

    await _staff(w, _press(f"done:{order.id}", note))
    assert _order(w, order).status == "delivered"
    assert w.telegram.to(CUSTOMER)[-1] == t("order_picked_up", "en", number=number, shop=STORE.name)

    await _staff(w, _press(f"done:{order.id}", note))
    assert _button_answers(w)[-1] == (f"Order #{number} was already picked up.", True)


@pytest.mark.anyio
async def test_pickup_steps_are_refused_unpaid_and_delivery_buttons_dont_fit():
    w = World()
    order = await placed_order(w)  # not paid
    await w.send_photo()
    alert = w.telegram.last_message_id(STAFF_CHAT)
    told = len(w.telegram.to(CUSTOMER))
    await _staff(w, _press(f"done:{order.id}", alert))
    answer, popup = _button_answers(w)[-1]
    assert popup and "isn't paid yet" in answer
    assert _order(w, order).status == "pending" and len(w.telegram.to(CUSTOMER)) == told

    w = World()
    order, note = await _paid_pickup(w)
    for step in ("ship", "back"):  # delivery buttons on a pickup order
        await _staff(w, _press(f"{step}:{order.id}", note))
        assert _button_answers(w)[-1] == ("This button doesn't fit this order.", True)
    assert _order(w, order).status == "confirmed"


@pytest.mark.anyio
async def test_the_pickup_message_ends_with_the_screenshot_request():
    w = World()
    await placed_order(w)
    assert w.last_text().endswith(t("after_paying", "en"))  # D78: the last thing they read


@pytest.mark.anyio
async def test_amharic_customer_hears_it_in_amharic():
    w = World(language="am")  # the customer chose Amharic (D29)
    await w.say("ኤር ፎርስ")
    for label in ("White", "42", "1", t("btn_continue", "am"), t("btn_delivery", "am")):
        await w.tap(label)
    await w.say("አበበ")
    await w.say("0911223344")
    await w.tap(t("btn_confirm", "am"))
    [order] = w.db.orders.values()
    alert = w.telegram.last_message_id(STAFF_CHAT)
    await _staff(w, _press(f"ship:{order.id}", alert))
    number = order_number(order.id)
    assert w.telegram.to(CUSTOMER)[-1].startswith(
        t("order_on_the_way_pay", "am", number=number, total="5,000 ብር"))
    await _staff(w, _press(f"done:{order.id}", alert))
    assert w.telegram.to(CUSTOMER)[-1] == t("order_delivered_paid", "am", number=number, shop=STORE.name)


@pytest.mark.anyio
async def test_another_stores_order_cant_be_moved():
    w = World()
    order, _ = await _delivery_order(w)
    other = STORE.model_copy(update={"id": uuid4()})
    result = await w.orchestrator.staff.advance_order(other, order.id, "ship")
    assert not result.ok and result.message == "Order not found."
    assert _order(w, order).status == "pending"


def test_the_buttons_follow_the_orders_stage():
    def order(**fields):
        return OrderWithItems(**{"id": uuid4(), "store_id": uuid4(), **fields})

    def steps(o, **kwargs):
        return [data.split(":")[0] for _, data in order_buttons(o, **kwargs)]

    assert steps(order(fulfillment_method="delivery")) == ["ship"]
    assert steps(order(fulfillment_method="delivery", status="out_for_delivery")) == ["done", "back"]
    assert steps(order(fulfillment_method="delivery", status="out_for_delivery", payment_status="paid")) == ["done"]
    assert steps(order(fulfillment_method="pickup")) == ["pay"]
    assert steps(order(fulfillment_method="pickup", status="confirmed", payment_status="paid")) == ["ready", "done"]
    assert steps(order(fulfillment_method="pickup", status="confirmed", payment_status="paid"),
                 ready_sent=True) == ["done"]
    assert steps(order(status="delivered", payment_status="paid")) == []
    assert steps(order(status="cancelled")) == []
    assert all(len(d) <= 64 for _, d in order_buttons(order(fulfillment_method="delivery",
                                                            status="out_for_delivery")))  # Telegram's limit


# --- 1. The morning summary (D70/D71) --------------------------------------------------------

RAW = {"revenue": 38500, "payments": 11, "orders_placed": 12, "telegram_orders": 9, "in_shop_sales": 3,
       "unpaid_orders": 2, "discount_total": 300,
       "top_products": [{"name": "Nike Air", "quantity": 5, "revenue": 25000}]}
LOW = [{"product_name": "Nike Air", "color": "Black", "size": "42", "stock": 1},
       {"product_name": "Leather bag", "color": "Brown", "size": None, "stock": 0}]


def summary_shop(**fields) -> Store:
    return Store(**{"id": uuid4(), "name": "Selam Shoes", "telegram_bot_token": TOKEN, "webhook_secret": "s",
                    "owner_telegram_id": 42, "staff_chat_id": -500, **fields})


def test_summary_says_how_yesterday_went():
    start, _ = yesterday(MORNING)
    text = summary_text(summary_shop(), RAW, LOW, start, "en")
    assert text.startswith("☀️ Selam Shoes · yesterday, Oct 3")
    for line in ("🛒 12 orders: 9 on Telegram, 3 in the shop", "💰 38,500 ETB received (11 payments)",
                 "⏳ 2 orders not paid yet", "🏷 Discounts in the shop: 300 ETB", "🏆 Best seller: Nike Air (5)",
                 "• Nike Air · Black · 42: 1 left", "• Leather bag · Brown: sold out"):
        assert line in text
    amharic = summary_text(summary_shop(), RAW, LOW, start, "am")
    assert "ትናንት" in amharic and "38,500 ብር" in amharic


def test_a_quiet_day_still_gets_a_summary():
    start, _ = yesterday(MORNING)
    text = summary_text(summary_shop(), {}, [], start, "en")
    assert "🛒 No orders yesterday." in text and "✅ Nothing is low on stock." in text
    assert "received" not in text and "not paid" not in text
    many = [{"product_name": f"P{i}", "stock": 1} for i in range(8)]
    assert "…and 3 more" in summary_text(summary_shop(), {}, many, start, "en")


def test_yesterday_is_an_addis_ababa_day():
    # 22:30 UTC on Oct 3 is already Oct 4 in Addis: yesterday is Oct 3.
    start, end = yesterday(datetime(2026, 10, 3, 22, 30, tzinfo=timezone.utc))
    assert (start.isoformat(), end - start) == ("2026-10-03T00:00:00+03:00", timedelta(days=1))


class SummaryDb:
    def __init__(self, stores):
        self.stores = stores
        self.sent_days = set()
        self.analytics = []

    async def stores_for_summary(self):
        return self.stores

    async def add_daily_summary(self, store_id, day):
        if (store_id, day) in self.sent_days:
            return False
        self.sent_days.add((store_id, day))
        return True

    async def store_analytics(self, store_id, start, end):
        self.analytics.append((store_id, start, end))
        return RAW

    async def low_stock(self, store_id, at_most=2, limit=50):
        return LOW


def summaries(db, refuse=()):
    sent = []

    def handler(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        if body["chat_id"] in refuse:
            return httpx.Response(403, json={"ok": False, "description": "Forbidden: bot was blocked by the user"})
        sent.append((body["chat_id"], body["text"]))
        return httpx.Response(200, json={"ok": True, "result": {"message_id": 1}})

    telegram = TelegramService(httpx.AsyncClient(base_url=TELEGRAM_API, transport=httpx.MockTransport(handler)))
    return DailySummaries(db, telegram), sent


@pytest.mark.anyio
async def test_each_owner_gets_it_once_a_day_in_their_language():
    english, amharic = summary_shop(daily_summary="en"), summary_shop(owner_telegram_id=43)
    db = SummaryDb([english, amharic])
    daily, sent = summaries(db)
    assert await daily.send_all(MORNING) == 2
    assert {chat for chat, _ in sent} == {42, 43}  # the owners' private chats only (D70)
    assert "yesterday" in dict(sent)[42] and "ትናንት" in dict(sent)[43]
    start, end = yesterday(MORNING)
    assert db.analytics[0] == (english.id, start, end)
    sent.clear()
    assert await daily.send_all(MORNING + timedelta(hours=3)) == 0 and sent == []  # once a day
    assert await daily.send_all(MORNING + timedelta(days=1)) == 2  # the next day, again


@pytest.mark.anyio
async def test_not_before_eight_and_a_blocked_bot_is_not_counted():
    db = SummaryDb([summary_shop()])
    daily, sent = summaries(db, refuse={42})
    assert await daily.maybe_send(datetime(2026, 10, 4, 7, 30, tzinfo=ADDIS)) == 0 and db.analytics == []
    assert await daily.maybe_send(MORNING) == 0  # tried, but the owner never pressed Start
    assert sent == [] and len(db.sent_days) == 1  # not retried every 15 minutes


def test_the_owner_turns_it_off_or_changes_the_language(world):
    db, _, client, _, _ = world
    response = client.put(mini.url("/settings"), headers=mini.headers(OWNER), json={"daily_summary": "off"})
    assert response.status_code == 200 and response.json()["daily_summary"] == "off"
    assert db.stores[STORE_A.id].daily_summary == "off"
    assert client.put(mini.url("/settings"), headers=mini.headers(OWNER),
                      json={"daily_summary": "fr"}).status_code == 422
    assert client.put(mini.url("/settings"), headers=mini.headers(MEMBER),
                      json={"daily_summary": "en"}).status_code == 403  # owners only


# --- 2. The export for the accountant (D72/D73, D79) ------------------------------------------

def test_gregorian_months_are_addis_months_and_show_their_ethiopian_days():
    period = month_bounds("2026-09", MORNING)
    assert (period.start.isoformat(), period.end.isoformat()) == ("2026-09-01T00:00:00+03:00",
                                                                  "2026-10-01T00:00:00+03:00")
    assert period.title == "September 2026" and period.slug == "2026-09"
    assert period.other == "ነሐሴ 26, 2018 – መስከረም 20, 2019 ዓ.ም."
    assert month_bounds("2025-12", MORNING).end.isoformat() == "2026-01-01T00:00:00+03:00"
    assert month_bounds("2026-10", MORNING).start.day == 1  # this month so far
    for bad in ("2026-13", "2026-9", "", "1999-01", "2026-11"):  # the last: not started yet
        with pytest.raises(ExportError):
            month_bounds(bad, MORNING)


def test_ethiopian_months_have_their_own_days():
    meskerem = month_bounds("2019-01", MORNING, "ethiopian")  # Meskerem 2019 = Sep 11 – Oct 10, 2026
    assert (meskerem.start.isoformat(), meskerem.end.isoformat()) == ("2026-09-11T00:00:00+03:00",
                                                                      "2026-10-11T00:00:00+03:00")
    assert meskerem.title == "መስከረም 2019 ዓ.ም. (Meskerem)" and meskerem.other == "Sep 11, 2026 – Oct 10, 2026"
    assert meskerem.slug == "Meskerem-2019-EC"
    nehase = month_bounds("2018-12", MORNING, "ethiopian")  # with Pagume, up to the new year
    assert (nehase.start.date(), nehase.end.date()) == (date(2026, 8, 7), date(2026, 9, 11))
    assert nehase.title.startswith("ነሐሴ + ጳጉሜ 2018")
    for bad in ("2019-13", "2019-00", "2011-01", "2019-02"):  # Pagume alone, too old, not started yet
        with pytest.raises(ExportError):
            month_bounds(bad, MORNING, "ethiopian")


def test_the_ethiopian_calendar():
    assert to_ethiopian(date(2026, 9, 11)) == (2019, 1, 1)  # new year
    assert to_ethiopian(date(2023, 9, 12)) == (2016, 1, 1)  # a day later before a Gregorian leap year
    assert to_ethiopian(date(2023, 9, 11)) == (2015, 13, 6)  # Pagume 6
    assert to_ethiopian(date(2027, 1, 7)) == (2019, 4, 29)  # Genna, Tahsas 29
    assert ethiopian_text(date(2026, 10, 6)) == "መስከረም 26, 2019"
    for ordinal in range(date(2020, 1, 1).toordinal(), date(2032, 1, 1).toordinal(), 7):
        day = date.fromordinal(ordinal)
        assert from_ethiopian(*to_ethiopian(day)) == day


def _row(status="confirmed", paid=True, channel="telegram", items=(), payments=None, **fields):
    return {"id": str(uuid4()), "status": status, "payment_status": "paid" if paid else "unpaid",
            "total_price": sum(i["price"] * i["quantity"] for i in items), "currency": "ETB",
            "fulfillment_method": "pickup", "contact_name": "Abebe", "contact_phone": "0911",
            "created_at": "2026-09-14T07:30:00+00:00", "channel": channel, "payment_method": None,
            "sold_by_name": None, "payments": payments if payments is not None else [], **fields,
            "order_items": [dict(i) for i in items]}


def _item(name, price, quantity=1, list_price=None, color="Black", size="42"):
    return {"quantity": quantity, "price": price, "list_price": list_price,
            "product_variants": {"color": color, "size": size, "products": {"name": name, "code": "P101"}}}


EXPORT_ROWS = [
    _row(items=[_item("Nike Air", 5000, 2)],
         payments=[{"amount": 10000, "method": "Telebirr", "paid_at": "x", "confirmed_by_name": "Sara"}]),
    _row(channel="in_shop", items=[_item("Leather bag", 900, 1, list_price=1000, color="Brown", size=None)],
         payment_method="Cash", sold_by_name="Hana",
         payments=[{"amount": 900, "method": "Cash", "paid_at": "x", "confirmed_by_name": "Hana"}]),
    _row(status="pending", paid=False, items=[_item("Nike Air", 5000)]),
    _row(status="cancelled", paid=False, items=[_item("Nike Air", 5000)]),
]


def _book(export):
    return load_workbook(io.BytesIO(export.data))


def test_the_workbook_has_every_sheet_in_both_calendars():
    export = build_export(STORE_A, EXPORT_ROWS, month_bounds("2026-09", MORNING), MORNING)
    assert export.filename == "Selam-Shoes-2026-09.xlsx" and export.revenue == Decimal("10900")
    book = _book(export)
    assert book.sheetnames == ["Summary", "By day", "By product", "Sales", "Orders"]

    rows = list(book["Summary"].iter_rows(values_only=True))
    assert rows[0][:2] == ("Selam Shoes", "September 2026")
    assert rows[1][1] == "ነሐሴ 26, 2018 – መስከረም 20, 2019 ዓ.ም."
    summary = {row[0]: row[1:] for row in rows if row[0]}
    assert summary["Money received (paid orders and sales)"][0] == 10900
    assert summary["Paid orders and sales"][0] == 2
    assert summary["Average sale"][0] == 5450
    assert summary["Discounts given in the shop"][0] == 100
    assert summary["Orders not paid"][0] == 1 and summary["Cancelled orders"][0] == 1
    assert summary["  Pickup (Telegram)"][0] == 2 and summary["  In the shop"][0] == 1
    assert summary["Telebirr"][:2] == (1, 10000) and summary["Cash"][:2] == (1, 900)

    days = list(book["By day"].iter_rows(values_only=True))
    assert days[0] == ("Date", "Ethiopian date", "Orders", "Paid", "Money received", "Discounts")
    assert len(days) == 31  # every day of September
    sep14 = next(d for d in days[1:] if d[0] == datetime(2026, 9, 14))
    assert sep14[1:] == ("መስከረም 4, 2019", 3, 2, 10900, 100)

    products = list(book["By product"].iter_rows(values_only=True))
    assert products[1] == ("Nike Air", "P101", 2, 10000, 0)  # the most money first
    assert products[2] == ("Leather bag", "P101", 1, 900, 100)

    sales = list(book["Sales"].iter_rows(values_only=True))
    assert sales[0][:9] == ("Date", "Ethiopian date", "Order", "Where", "Product", "Code", "Color", "Size",
                            "Quantity")
    assert len(sales) == 3  # only paid, not cancelled
    nike, bag = sales[1], sales[2]
    assert nike[:2] == (datetime(2026, 9, 14, 10, 30), "መስከረም 4, 2019")  # Addis time, both calendars
    assert nike[4:] == ("Nike Air", "P101", "Black", "42", 2, 5000, 5000, 0, 10000, "Telebirr", "Sara")
    assert bag[3] == "In the shop" and bag[9:] == (1000, 900, 100, 900, "Cash", "Hana")

    orders = list(book["Orders"].iter_rows(values_only=True))
    assert len(orders) == 5  # every order, with its stage
    assert [o[7] for o in orders[1:]] == ["Paid", "Paid", "Waiting for payment", "Cancelled"]


def test_an_ethiopian_month_file():
    export = build_export(STORE_A, EXPORT_ROWS, month_bounds("2019-01", MORNING, "ethiopian"), MORNING)
    assert export.filename == "Selam-Shoes-Meskerem-2019-EC.xlsx"
    book = _book(export)
    assert next(book["Summary"].iter_rows(values_only=True))[:2] == ("Selam Shoes", "መስከረም 2019 ዓ.ም. (Meskerem)")
    days = list(book["By day"].iter_rows(values_only=True))
    assert days[1][:2] == (datetime(2026, 9, 11), "መስከረም 1, 2019")
    assert days[-1][:2] == (datetime(2026, 10, 4), "መስከረም 24, 2019")  # up to today


def test_the_columns_use_the_shops_own_words():
    phones = STORE_A.model_copy(update={"shop_type": "electronics"})
    sales = next(_book(build_export(phones, [], month_bounds("2026-09", MORNING), MORNING))["Sales"]
                 .iter_rows(values_only=True))
    assert sales[6:8] == ("Color", "Storage")


def test_the_owner_gets_the_file_from_the_bot(world):
    db, telegram, client, _, _ = world
    asked = []

    async def orders_for_export(store_id, start, end):
        asked.append((store_id, start, end))
        return EXPORT_ROWS

    db.orders_for_export = orders_for_export
    response = client.post(mini.url("/orders/export"), headers=mini.headers(OWNER), json={"month": "2026-09"})
    assert response.status_code == 200, response.text
    assert response.json()["file"] == "Selam-Shoes-2026-09.xlsx" and response.json()["orders"] == 4
    assert [store_id for store_id, _, _ in asked] == [STORE_A.id]  # only this shop's orders
    token, form = telegram.documents[-1]
    assert token == mini.TOKEN_A  # the shop's own bot...
    assert b'name="chat_id"\r\n\r\n1\r\n' in form  # ...to the owner who asked, in private
    assert b"Selam-Shoes-2026-09.xlsx" in form

    # The same, by Ethiopian month (D79).
    response = client.post(mini.url("/orders/export"), headers=mini.headers(OWNER),
                           json={"month": "2019-01", "calendar": "ethiopian"})
    assert response.status_code == 200 and response.json()["file"] == "Selam-Shoes-Meskerem-2019-EC.xlsx"
    assert asked[-1][1].isoformat() == "2026-09-11T00:00:00+03:00"


def test_export_is_for_owners_and_needs_the_bot_started(world):
    db, telegram, client, _, _ = world

    async def orders_for_export(store_id, start, end):
        return []

    db.orders_for_export = orders_for_export
    path = mini.url("/orders/export")
    assert client.post(path, headers=mini.headers(MEMBER), json={"month": "2026-09"}).status_code == 403
    assert client.post(path, headers=mini.headers(OWNER), json={"month": "2026-9"}).status_code == 422
    assert client.post(path, headers=mini.headers(OWNER), json={"month": "2099-01"}).status_code == 422
    assert client.post(path, headers=mini.headers(OWNER),
                       json={"month": "2019-13", "calendar": "ethiopian"}).status_code == 422
    assert client.post(path, headers=mini.headers(OWNER),
                       json={"month": "2026-09", "calendar": "julian"}).status_code == 422
    telegram.refuse_documents = True
    response = client.post(path, headers=mini.headers(OWNER), json={"month": "2026-09"})
    assert response.status_code == 409 and "@selam_bot" in response.json()["detail"]
    # Another store's owner can't export this store.
    assert client.post(path, headers=mini.headers(mini.STRANGER, mini.TOKEN_B),
                       json={"month": "2026-09"}).status_code in (401, 403)
